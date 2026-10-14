from pathlib import Path
import csv
import json
import re
import math
import hashlib
import subprocess
from collections import Counter, defaultdict
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics import renderPDF
from reportlab.lib import colors
from equations import equation

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
POPPLER = Path(r'C:\Users\study\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin')
sources = []

def read_csv(path):
    path = ROOT / path
    sources.append(path)
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def read_json(path):
    path = ROOT / path
    sources.append(path)
    return json.loads(path.read_text(encoding='utf-8'))

def best(rows):
    return min(rows, key=lambda r: float(r['val_loss']))

main = {s: read_csv(f'outputs/char_token/20261013_161958/{s}_train_epochs.csv') for s in ['sub-01','sub-02']}
noise = {s: read_csv(f'outputs/char_token_noise/20261013_175435/{s}_sub-01_train_epochs.csv') for s in ['eeg','noise']}
pool = {s: read_csv(f'outputs/char_token_pooling/20261013_200409/{s}_sub-01_train_epochs.csv') for s in ['attn','mean']}
probes = {(p,x): read_json(f'outputs/char_token_pooling/20261013_200409/probe_{p}_sub-01_{x}.json') for p in ['attn','mean'] for x in ['real','noise']}
folds = read_csv('outputs/char_token_day_holdout/20261013_221128/four_fold_summary.csv')
cross = [read_json(p.relative_to(ROOT)) for p in sorted((ROOT/'outputs/char_token_cross_subject').glob('*/*metrics.json'))]
meta_path = ROOT/'datasets/preprocessed/all_text.tsv'
sources.append(meta_path)
with meta_path.open(encoding='utf-8-sig', newline='') as f:
    meta = list(csv.DictReader(f, delimiter='\t'))
meta = [r for r in meta if int(r['sample']) == 250]
groups = defaultdict(list)
for row in meta:
    groups[(row['subject'], 'para1' if int(row['day']) <= 4 else 'para2')].append(row)

def pct(value):
    return f'{100*float(value):.2f}'

def metrics_row(label, r):
    return [label, str(r['epoch']), f"{float(r['val_loss']):.4f}", *[pct(r[k]) for k in ['eeg2text_top1','eeg2text_top5','text2eeg_top1','text2eeg_top5']]]

tables = {}
tables['data'] = (['被试','语料','试次数','唯一句子','字符范围'], [[s,p,str(len(v)),str(len({r['sentence'] for r in v})),f"{min(int(r['n_chars']) for r in v)}～{max(int(r['n_chars']) for r in v)}"] for (s,p),v in sorted(groups.items())])
metric_head = ['设置','轮次','验证损失','E→T@1 %','E→T@5 %','T→E@1 %','T→E@5 %']
tables['main'] = (metric_head,[metrics_row(s,best(r)) for s,r in main.items()])
tables['noise'] = (metric_head,[metrics_row(s,best(r)) for s,r in noise.items()])
tables['pool'] = (metric_head,[metrics_row(s,best(r)) for s,r in pool.items()])
tables['probe'] = (['池化','输入','R²','MAE','置乱标签 R²','检索 Top5 %'], [[p,'真实 EEG' if x=='real' else '噪声',f"{v['length_probe']['val_r2']:.4f}",f"{v['length_probe']['val_mae']:.4f}",f"{v['length_probe']['null_r2']:.4f}",pct(v['stratified_retrieval']['overall_top5'])] for (p,x),v in probes.items()])
tables['fold'] = (['留出天','验证损失','E→T@1 %','E→T@5 %','T→E@1 %','T→E@5 %'], [[r['val_day'] if r['fold'] not in ['mean','std'] else ('均值' if r['fold']=='mean' else '标准差'),f"{float(r['val_loss']):.4f}",*[pct(r[k]) for k in ['eeg2text_top1','eeg2text_top5','text2eeg_top1','text2eeg_top5']]] for r in folds])
tables['cross'] = (['源→目标','验证损失','E→T@1 %','E→T@5 %','T→E@1 %','T→E@5 %'], [[r['source'].replace('sub-','S')+'→'+r['target'].replace('sub-','S'),f"{r['val_loss']:.4f}",*[pct(r[k]) for k in ['eeg2text_top1','eeg2text_top5','text2eeg_top1','text2eeg_top5']]] for r in cross])
tables['length'] = (['token 数','验证试次数','注意力 Top5 %','均值 Top5 %'], [[str(a['n_tok']),str(a['n_trial']),pct(a['top5']),pct(b['top5'])] for a,b in zip(probes[('attn','real')]['stratified_retrieval']['by_n_tok'],probes[('mean','real')]['stratified_retrieval']['by_n_tok'])])
tables['config'] = (['项目','主实验和对照','跨天实验'],[['采样率与 token 时长','250 Hz 和 0.4 s','相同'],['编码器宽度和层数','128 维和 2 层','相同'],['注意力头数与前馈宽度','4 头和 256 维','相同'],['文本向量维数','768','相同'],['每批句子数和重复数','32 × 4','32 × 3'],['训练轮数','30','20'],['初始学习率和权重衰减','均为 0.001','相同'],['温度和随机种子','0.2 和 0','相同'],['优化器和调度','AdamW 和余弦退火','相同'],['模型选择','最小验证损失','相同']])

def save_chart(drawing, name):
    pdf = ASSETS/(name+'.pdf')
    renderPDF.drawToFile(drawing, str(pdf))
    subprocess.run([str(POPPLER/'pdftoppm.exe'),'-png','-r','180','-singlefile',str(pdf),str(ASSETS/name)],check=True,capture_output=True)

def line_chart(name, series, ylabel, ymin, ymax):
    d = Drawing(520,270)
    lp = LinePlot()
    lp.x,lp.y,lp.width,lp.height = 58,49,435,182
    lp.data = [data for _,data in series]
    lp.xValueAxis.valueMin=1
    lp.xValueAxis.valueMax=max(x for _,data in series for x,_ in data)
    lp.xValueAxis.valueStep=5
    lp.yValueAxis.valueMin=ymin
    lp.yValueAxis.valueMax=ymax
    lp.yValueAxis.valueStep=(ymax-ymin)/4
    lp.yValueAxis.visibleGrid=1
    lp.yValueAxis.gridStrokeColor=colors.HexColor('#dddddd')
    palette=['#176b8b','#c04d40','#458447','#76518d']
    for i,(label,_) in enumerate(series):
        lp.lines[i].strokeColor=colors.HexColor(palette[i])
        lp.lines[i].strokeWidth=1.8
        d.add(String(58+i*(435/len(series)),255,label,fontSize=9,fillColor=colors.HexColor(palette[i])))
    d.add(lp)
    d.add(String(262,12,'Epoch',fontSize=10,textAnchor='middle'))
    d.add(String(58,236,ylabel,fontSize=10))
    save_chart(d,name)

line_chart('main_loss',[(s,[(int(r['epoch']),float(r['val_loss'])) for r in rows]) for s,rows in main.items()], 'Validation loss',6.2,6.9)
line_chart('main_top5',[(s,[(int(r['epoch']),100*float(r['eeg2text_top5'])) for r in rows]) for s,rows in main.items()]+[('Chance',[(1,100*5/225),(30,100*5/225)])], 'EEG to text Top5 (%)',0,12)
line_chart('noise_loss',[(s,[(int(r['epoch']),float(r['val_loss'])) for r in rows]) for s,rows in noise.items()], 'Validation loss',6.2,6.9)
line_chart('pool_loss',[(s,[(int(r['epoch']),float(r['val_loss'])) for r in rows]) for s,rows in pool.items()], 'Validation loss',6.2,6.9)
fold_series=[]
for i in range(1,5):
    rows=read_csv(f'outputs/char_token_day_holdout/20261013_221128/fold{i}_val-day{i:02d}/sub-01_train_epochs.csv')
    fold_series.append((f'Held-out day {i}',[(int(r['epoch']),float(r['val_loss'])) for r in rows]))
line_chart('fold_loss',fold_series,'Validation loss',7.0, max(7.5,math.ceil(max(y for _,data in fold_series for _,y in data)*10)/10))

def chart_bars(name, categories, vals, ylabel, ymax):
    d=Drawing(520,260)
    bc=VerticalBarChart()
    bc.x,bc.y,bc.width,bc.height=55,50,435,170
    bc.data=[vals]
    bc.categoryAxis.categoryNames=categories
    bc.categoryAxis.labels.fontSize=9
    bc.valueAxis.valueMin=0
    bc.valueAxis.valueMax=ymax
    bc.valueAxis.valueStep=ymax/5
    bc.bars[0].fillColor=colors.HexColor('#176b8b')
    bc.barLabels.nudge=7
    bc.barLabelFormat='%.3f'
    d.add(bc)
    d.add(String(55,238,ylabel,fontSize=11))
    save_chart(d,name)

chart_bars('probe_r2',['Attn EEG','Attn noise','Mean EEG','Mean noise'],[v['length_probe']['val_r2'] for v in probes.values()],'Length probe R-squared',0.6)
chart_bars('fold_top5',[f'Day {i}' for i in range(1,5)]+['Chance'],[100*float(r['eeg2text_top5']) for r in folds[:4]]+[100*5/1124],'EEG to text Top5 (%)',0.7)

facts={'tables':tables,'metadata_rows':len(meta),'sources':[]}
for path in sorted(set(sources)):
    facts['sources'].append({'path':str(path.relative_to(ROOT)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
(OUT/'evidence.json').write_text(json.dumps(facts,ensure_ascii=False,indent=2),encoding='utf-8')

text=(OUT/'thesis_source.md').read_text(encoding='utf-8')
for name,(heads,rows) in tables.items():
    table='| '+' | '.join(heads)+' |\n| '+' | '.join(['---']*len(heads))+' |\n'+'\n'.join('| '+' | '.join(row)+' |' for row in rows)
    text=text.replace('{{TABLE:'+name+'}}',table)
text=text.replace('{{META_ROWS}}',str(len(meta)))
outline='\n'.join(('  ' if line.startswith('## ') else '')+'- '+line.lstrip('# ') for line in text.splitlines() if line.startswith(('# ','## ')))
(OUT/'毕业论文完整初稿.md').write_text(text.replace('{{TOC}}','# 目录\n\n'+outline),encoding='utf-8')

doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.top_margin,sec.bottom_margin=Cm(2.5),Cm(2.5)
sec.left_margin,sec.right_margin=Cm(2.8),Cm(2.5)
sec.header_distance,sec.footer_distance=Cm(1.2),Cm(1.2)
styles=doc.styles
for border in styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=styles[name]
    style.font.name='Times New Roman'
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '黑体')
normal=styles['Normal']
normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(3)
normal.paragraph_format.first_line_indent=Pt(24)
normal.paragraph_format.widow_control=True
for name,size in [('Title',24),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
    styles[name].font.size=Pt(size)
    styles[name].paragraph_format.first_line_indent=Pt(0)
    styles[name].paragraph_format.space_before=Pt(12)
    styles[name].paragraph_format.space_after=Pt(8)
    styles[name].paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.page_break_before=True
styles['Subtitle'].font.italic=False
for name in ['Title','Subtitle']:
    for spacing in styles[name].element.xpath('.//w:rPr/w:spacing'):
        spacing.getparent().remove(spacing)
for level in [1,2]:
    name=f'TOC {level}'
    style=styles[name] if name in styles else styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    style.base_style=normal
    style.font.size=Pt(11)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.line_spacing=1.15
    style.paragraph_format.space_after=Pt(2)
    style.paragraph_format.first_line_indent=Pt(0)
    style.paragraph_format.left_indent=Cm(0 if level==1 else 0.6)
    style.paragraph_format.keep_with_next=False
    style.paragraph_format.page_break_before=False
styles['Caption'].font.size=Pt(10.5)
styles['Caption'].paragraph_format.first_line_indent=Pt(0)
styles['Caption'].paragraph_format.line_spacing=1.1
styles['Caption'].paragraph_format.space_after=Pt(8)

def field(paragraph, code, cached='1'):
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=code
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for el in [begin,instr,sep]: paragraph.add_run()._r.append(el)
    paragraph.add_run(cached)
    paragraph.add_run()._r.append(end)

sec.different_first_page_header_footer=True
hp=sec.header.paragraphs[0]
hp.text='字符级脑电与中文文本对比学习及泛化能力研究'
hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.first_line_indent=Pt(0)
for run in hp.runs:run.font.size=Pt(9)
fp=sec.footer.paragraphs[0]
fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent=Pt(0)
field(fp,' PAGE ')

doc.add_paragraph('硕士学位论文',style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n')
p=doc.add_paragraph('字符级脑电与中文文本\n对比学习及泛化能力研究',style='Title')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Character-Level EEG and Chinese Text\nContrastive Learning and Generalization',style='Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')
for item in ['学校名称：____________________','学院名称：____________________','学科专业：____________________','研究方向：脑机接口与机器学习','研究生姓名：__________________','学号：________________________','指导教师：____________________','完成日期：2026 年 10 月']:
    p=doc.add_paragraph(item)
    p.paragraph_format.first_line_indent=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def add_table(lines):
    parsed=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    parsed=[r for r in parsed if not all(re.fullmatch(r'[-: ]+',c or '-') for c in r)]
    table=doc.add_table(rows=1,cols=len(parsed[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    width=15.7/len(parsed[0])
    for col in table.columns:col.width=Cm(width)
    for ri,row in enumerate(parsed):
        cells=table.rows[0].cells if ri==0 else table.add_row().cells
        for ci,value in enumerate(row):
            cell=cells[ci]
            cell.text=value
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge);m.set(qn('w:w'),'80');m.set(qn('w:type'),'dxa');margins.append(m)
            tcpr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF0');tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=1.1
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.keep_with_next=ri<len(parsed)-1
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size=Pt(9 if len(row)>5 else 10)
                    run.bold=ri==0
        trpr=table.rows[ri]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=text.splitlines()
i=0
current_section=''
while i<len(lines):
    line=lines[i].strip()
    i+=1
    if not line:continue
    if line=='{{TOC}}':
        p=doc.add_paragraph('目录',style='Title')
        p.paragraph_format.page_break_before=True
        p=doc.add_paragraph()
        p.paragraph_format.first_line_indent=Pt(0)
        field(p,' TOC \\o "1-2" \\h \\z \\u ','目录')
        continue
    if line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i].strip());i+=1
        add_table(block);continue
    if line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(OUT/m[2]),width=Cm(15.5))
        p.paragraph_format.keep_with_next=True
        doc.add_paragraph(m[1],style='Caption').alignment=WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('# '):
        current_section=line[2:]
        doc.add_paragraph(line[2:],style='Heading 1');continue
    if line.startswith('## '):
        doc.add_paragraph(line[3:],style='Heading 2');continue
    if line.startswith('### '):
        doc.add_paragraph(line[4:],style='Heading 3');continue
    if line.startswith('表 '):
        p=doc.add_paragraph(line,style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;continue
    if line.startswith('公式：'):
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing=1.2
        p.paragraph_format.space_before=Pt(6)
        p.paragraph_format.space_after=Pt(6)
        number=re.search(r'（([0-9]+-[0-9]+)）',line).group(1)
        p._p.append(equation(number))
        p.add_run('    ('+number+')')
        continue
    p=doc.add_paragraph(line)
    if current_section in ['摘要','Abstract','ABSTRACT']:
        p.paragraph_format.line_spacing=1.2
        p.paragraph_format.first_line_indent=Pt(22)
        for run in p.runs:run.font.size=Pt(11)
    if line.startswith('[') or line.startswith('python ') or line.startswith('uv run'):
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.line_spacing=1.2
        for run in p.runs:run.font.size=Pt(10.5)

doc.core_properties.title='字符级脑电与中文文本对比学习及泛化能力研究'
doc.core_properties.subject='硕士学位论文'
doc.core_properties.author=''
doc.core_properties.keywords='EEG;中文文本;对比学习;跨天泛化;长度探针'
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
path=OUT/'毕业论文完整初稿.docx'
doc.save(path)
print(json.dumps({'docx':str(path),'chinese_characters':len(re.findall(r'[\u4e00-\u9fff]',text)),'characters':len(text),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
